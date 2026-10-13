# app/api/routes/gmp_dashboard.py
# GMP counterpart of app/api/routes/dashboard.py — same shape (received /
# completed / on-process / summary / chart / recent-applications / detail),
# scoped to GMPRecord / GMPApplicationLogs instead of MainDB / ApplicationLogs.

from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Optional, Literal
from datetime import date, datetime
import io

from app.db.session import get_db
from app.core.deps import get_current_active_user
from app.crud import gmp_dashboard as crud_dashboard
from app.crud import gmp_dashboard_chart as crud_chart
from app.crud import gmp_dashboard_recent as crud_recent
from app.crud import gmp_dashboard_detail as crud_detail
from app.crud import gmp_post_eval_status as crud_post_eval
from app.schemas.dashboard import StatResponse, CombinedStatsResponse
from app.schemas.dashboard_chart import ChartResponse
from app.schemas.recent_applications import RecentApplicationsResponse
from app.schemas.gmp_dashboard_detail import GMPMetricDetailResponse
from app.schemas.gmp_post_eval_status import (
    PostEvalStatusSummaryResponse,
    PostEvalStatusListResponse,
)
from app.models.user import User

router = APIRouter(
    prefix="/api/dashboard/gmp/stats",
    tags=["GMP Dashboard Stats"],
)


def _effective_user(
    current_user,
    impersonate: Optional[int],
    db: Session = None,
) -> tuple[str, Optional[int]]:
    """(username, user_id) — the impersonated target's, or the current user's
    own. Callers should pass BOTH into the crud layer's _assignee_match
    (user_id OR user_name) rather than username alone, which silently drops
    every log recorded before the user's last rename."""
    if impersonate and db and current_user.role in ("Admin", "SuperAdmin"):
        user = db.query(User).filter(User.id == impersonate).first()
        if not user:
            raise HTTPException(status_code=404, detail="User not found")
        return user.username, user.id
    return current_user.username, current_user.id


def _common_params(
    date_from: Optional[date] = Query(None, description="Start date filter  YYYY-MM-DD"),
    date_to: Optional[date] = Query(None, description="End date filter    YYYY-MM-DD"),
    impersonate: Optional[int] = Query(
        None, description="Admin only: user_id of target user"
    ),
):
    return {"date_from": date_from, "date_to": date_to, "impersonate": impersonate}


@router.get(
    "/received",
    response_model=StatResponse,
    summary="Total GMP applications received by the current user",
)
def get_received(
    params: dict = Depends(_common_params),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, params["impersonate"], db)
    value = crud_dashboard.get_total_received(
        db, username, user_id, params["date_from"], params["date_to"]
    )
    return StatResponse(
        label="Total Received",
        value=value,
        username=username,
        date_from=params["date_from"],
        date_to=params["date_to"],
    )


@router.get(
    "/completed",
    response_model=StatResponse,
    summary="Total completed GMP applications for the current user",
)
def get_completed(
    params: dict = Depends(_common_params),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, params["impersonate"], db)
    value = crud_dashboard.get_total_completed(
        db, username, user_id, params["date_from"], params["date_to"]
    )
    return StatResponse(
        label="Completed",
        value=value,
        username=username,
        date_from=params["date_from"],
        date_to=params["date_to"],
    )


@router.get(
    "/on-process",
    response_model=StatResponse,
    summary="Total on-process GMP applications for the current user",
)
def get_on_process(
    params: dict = Depends(_common_params),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, params["impersonate"], db)
    value = crud_dashboard.get_total_on_process(
        db, username, user_id, params["date_from"], params["date_to"]
    )
    return StatResponse(
        label="On Process",
        value=value,
        username=username,
        date_from=params["date_from"],
        date_to=params["date_to"],
    )


@router.get(
    "/summary",
    response_model=CombinedStatsResponse,
    summary="All 3 GMP KPI stats in one request",
)
def get_summary(
    params: dict = Depends(_common_params),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, params["impersonate"], db)
    stats = crud_dashboard.get_stats_summary(
        db, username, user_id, params["date_from"], params["date_to"]
    )
    return CombinedStatsResponse(
        username=username,
        date_from=params["date_from"],
        date_to=params["date_to"],
        **stats,
    )


@router.get(
    "/chart",
    response_model=ChartResponse,
    summary="Time-series data for the GMP Insights chart and data table",
)
def get_chart(
    breakdown: Literal["day", "month", "year"] = Query(
        "day",
        description="Aggregation granularity",
    ),
    date_from: Optional[date] = Query(None, description="Inclusive start date  YYYY-MM-DD"),
    date_to: Optional[date] = Query(None, description="Inclusive end date    YYYY-MM-DD"),
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail="date_from must be earlier than or equal to date_to.",
        )

    username, user_id = _effective_user(current_user, impersonate, db)

    try:
        return crud_chart.get_chart_data(
            db=db,
            username=username,
            breakdown=breakdown,
            user_id=user_id,
            date_from=date_from,
            date_to=date_to,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))


@router.get(
    "/recent-applications",
    response_model=RecentApplicationsResponse,
    summary="Most recent GMP application log entries for the current user",
)
def get_recent_applications(
    page: int = Query(default=1, ge=1, description="1-based page number"),
    page_size: int = Query(default=10, ge=1, le=50, description="Rows per page (max 50)"),
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, impersonate, db)

    data = crud_recent.get_recent_applications(
        db=db,
        username=username,
        user_id=user_id,
        page=page,
        page_size=page_size,
    )

    return RecentApplicationsResponse(
        data=data["rows"],
        count=len(data["rows"]),
        total=data["total"],
        total_pages=data["total_pages"],
        page=data["page"],
        username=username,
    )


@router.get(
    "/detail",
    response_model=GMPMetricDetailResponse,
    summary="Paginated GMP application log rows for a specific KPI metric",
)
def get_metric_detail(
    metric: str = Query(..., description="KPI to drill into: received | completed | on_process"),
    date_from: Optional[date] = Query(None, description="Inclusive start date  YYYY-MM-DD  (filters start_date)"),
    date_to: Optional[date] = Query(None, description="Inclusive end date    YYYY-MM-DD  (filters start_date)"),
    accomplished_date_from: Optional[date] = Query(None, description="Filter rows where end_date >= this date"),
    accomplished_date_to: Optional[date] = Query(None, description="Filter rows where end_date <= this date"),
    app_step: Optional[str] = Query(None, description="Filter rows by app_step"),
    dtn: Optional[str] = Query(None, description="Filter by DTN (partial match)"),
    sort_by: Optional[str] = Query(None, description="Column to sort by. Currently supported: 'dtn'"),
    sort_dir: Literal["asc", "desc"] = Query("asc", description="Sort direction when sort_by is provided"),
    page: int = Query(default=1, ge=1, description="1-based page number"),
    page_size: int = Query(default=10, ge=1, le=500, description="Rows per page (max 500)"),
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    if metric not in ("received", "completed", "on_process"):
        raise HTTPException(
            status_code=422,
            detail="metric must be one of: received, completed, on_process",
        )

    if date_from and date_to and date_from > date_to:
        raise HTTPException(
            status_code=422,
            detail="date_from must be earlier than or equal to date_to.",
        )

    if (
        accomplished_date_from
        and accomplished_date_to
        and accomplished_date_from > accomplished_date_to
    ):
        raise HTTPException(
            status_code=422,
            detail="accomplished_date_from must be earlier than or equal to accomplished_date_to.",
        )

    username, user_id = _effective_user(current_user, impersonate, db)

    try:
        return crud_detail.get_metric_detail(
            db=db,
            username=username,
            user_id=user_id,
            metric=metric,
            date_from=date_from,
            date_to=date_to,
            accomplished_date_from=accomplished_date_from,
            accomplished_date_to=accomplished_date_to,
            app_step=app_step,
            dtn=dtn,
            sort_by=sort_by,
            sort_dir=sort_dir,
            page=page,
            page_size=page_size,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))


# ── Post-Evaluation Status (FGMP Dashboard only — no CPR counterpart) ─────────
# Applications the current user has finished their own part on (a COMPLETED
# step they were the assignee of) that aren't released yet — i.e. still
# actively at some step (GMP_CURRENT_STEP is not null). Powers the
# "Post-Evaluation Status" card: a labelled count breakdown for the header,
# plus a paginated row list for its "See all". Each breakdown entry is tagged
# "returned" (sitting back on the step this user last completed — theirs to
# act on) or "forward" (moved on to someone else), and its label doubles as
# the `bucket` the list endpoint filters on.
@router.get(
    "/post-eval-status/summary",
    response_model=PostEvalStatusSummaryResponse,
    summary="Grouped label counts for the Post-Evaluation Status card's header chips",
)
def get_post_eval_status_summary(
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, impersonate, db)
    breakdown = crud_post_eval.get_step_breakdown(db, username, user_id)
    return PostEvalStatusSummaryResponse(
        breakdown=breakdown,
        total=sum(b["count"] for b in breakdown),
    )


@router.get(
    "/post-eval-status",
    response_model=PostEvalStatusListResponse,
    summary="Paginated rows for the Post-Evaluation Status card / its 'See all'",
)
def get_post_eval_status(
    page: int = Query(default=1, ge=1, description="1-based page number"),
    page_size: int = Query(default=10, ge=1, le=50, description="Rows per page (max 50)"),
    bucket: Optional[str] = Query(
        None,
        description=(
            "Filter to one header chip: 'all' (default), 'returned' for every "
            "return at once, or the exact label from /post-eval-status/summary "
            "(e.g. 'Returned by Checker', 'Checker')"
        ),
    ),
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    username, user_id = _effective_user(current_user, impersonate, db)
    data = crud_post_eval.get_rows(
        db, username, user_id, page=page, page_size=page_size, bucket=bucket
    )
    return PostEvalStatusListResponse(
        data=data["rows"],
        total=data["total"],
        total_pages=data["total_pages"],
        page=data["page"],
    )


@router.get(
    "/post-eval-status/export",
    summary="Export the Post-Evaluation Status list to .xlsx",
    response_class=StreamingResponse,
)
def export_post_eval_status(
    buckets: Optional[List[str]] = Query(
        None,
        description=(
            "Repeat per chip to export, e.g. ?buckets=Checker&buckets=Returned by "
            "Checker. 'all' (or omitting this) exports everything. The values are "
            "the labels from /post-eval-status/summary."
        ),
    ),
    impersonate: Optional[int] = Query(None, description="Admin only — user_id of target user"),
    db: Session = Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    """Mirrors the on-screen table's columns, for the union of the chosen
    chips. Styling follows /api/gmp/export-filtered so the two exports look
    like they came from the same system."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    username, user_id = _effective_user(current_user, impersonate, db)
    rows = crud_post_eval.get_export_rows(db, username, user_id, buckets=buckets)
    if not rows:
        raise HTTPException(status_code=404, detail="Nothing to export for that selection")

    HEADERS = ["DTN", "Establishment", "Your Step", "Completed", "Where It Is Now", "Group"]
    WIDTHS = [20, 42, 20, 14, 28, 16]
    GROUP_LABELS = {
        crud_post_eval.GROUP_RETURNED: "Returned to me",
        crud_post_eval.GROUP_FORWARD: "Moved forward",
    }

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Post-Evaluation Status")

    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_align = Alignment(horizontal="center", vertical="center")

    header_cells = []
    for title in HEADERS:
        cell = WriteOnlyCell(ws, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        header_cells.append(cell)
    ws.append(header_cells)

    for r in rows:
        completed = r["completed_date"]
        ws.append([
            r["dtn"],
            r["lto_company"],
            r["your_step"],
            completed.strftime("%Y-%m-%d") if completed else None,
            r["label"],
            GROUP_LABELS.get(r["group"], r["group"]),
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"post_evaluation_status_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Total-Records": str(len(rows)),
        },
    )
