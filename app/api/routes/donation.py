# ============================================================
# FILE: app/api/routes/donation.py
# ============================================================
"""Donation Database routes.

The "Download Template" endpoint mirrors the header design of the
original CRR "Donation Database" encoding workbook (light-blue for the
CRR-filled DTN/date columns, pale-yellow for STATUS, light-gray for
Donation Registration No., bold Arial Narrow on the rest — thin
borders, centered + wrapped, frozen header row).
"""
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from typing import List, Optional
from datetime import datetime, date
import io

from app.core.deps import get_current_active_user
from app.db.session import get_db
from app.models.user import User
from app.schemas.donation import (
    DonationCreate,
    DonationUpdate,
    DonationResponse,
    DonationChangeLogResponse,
    DonationUploadResult,
)
from app.crud import donation as crud

router = APIRouter(
    prefix="/api/donation",
    tags=["Donation Database"],
    dependencies=[Depends(get_current_active_user)],
)

# (header label, column width, fill color, field key) — order matches the
# original workbook AND app/models/donation.py's column order. The field
# key is how an uploaded template's columns get mapped onto the model when
# importing (by position, not by header text).
TEMPLATE_COLUMNS = [
    ("Letter DTN\n(To be accomplished by CRR)", 26, "DAEEF3", "letter_dtn"),
    ("Date Received By Center\n(DD-MMM-YYYY)\n(To be accomplished by CRR)", 23, "DAEEF3", "date_received"),
    ("Date Received by Evaluator\n(DD-MMM-YYYY)", 22, None, "date_received_by_evaluator"),
    ("Donor", 29, None, "donor"),
    ("Donee/Recipient", 29, None, "donee"),
    ("Registration DTN", 23, None, "registration_dtn"),
    ("Product Name", 33, None, "product_name"),
    ("Packaging", 31, None, "packaging"),
    ("Manufacturer", 40, None, "manufacturer"),
    ("Batch/Lot No.", 24, None, "batch_lot_no"),
    ("Expiration Date\n(DD-MMM-YYYY)", 19, None, "expiration_date"),
    ("Total Quantity", 22, None, "total_quantity"),
    ("Validity (Expired on)\n(DD-MMM-YYYY)", 20, None, "validity"),
    ("Date Issue\n(DD-MMM-YYYY)", 20, None, "date_issued"),
    ("Evaluator", 13, None, "evaluator"),
    ("STATUS\n(Approved or Disapproved)", 19, "FFFFCC", "status"),
    ("Donation Registration No.", 20, "F2F2F2", "donation_reg_no"),
    ("Date Forwarded to Checker\n(DD-MMM-YYYY)", 19, None, "date_forwarded_to_checker"),
    ("Date Released from CDRR\n(DD-MMM-YYYY)\n(To be accomplished by CRR)", 19, "DAEEF3", "date_released"),
    ("REMARKS", 30, None, "remarks"),
]

# The columns that hold a date — accept whatever format the uploaded file
# has them in (a real Excel date, an Excel day-serial number, or free text
# like "28-May-2026") but always store them as YYYY-MM-DD, so a date typed
# any which way in Excel ends up consistent in the database (and matches
# what the Filters modal's <input type="date"> compares against).
DATE_FIELD_KEYS = {
    "date_received",
    "date_received_by_evaluator",
    "expiration_date",
    "validity",
    "date_issued",
    "date_forwarded_to_checker",
    "date_released",
}

_DATE_TEXT_FORMATS = (
    "%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y",
    "%B %d, %Y", "%b %d, %Y", "%d-%m-%Y", "%m-%d-%Y",
)


def _parse_import_date(value):
    """Best-effort: turn one date cell from an uploaded Excel file into
    YYYY-MM-DD before it's stored. Accepts a native datetime/date
    (what openpyxl gives back for a date-formatted cell), an Excel
    day-serial number (a date-formatted cell that lost its number format
    on save/re-save), or free text in a handful of common layouts. Falls
    back to the original text untouched if none of that matches, so a
    value is never silently dropped just because it doesn't parse."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:  # noqa: BLE001
            return str(value).strip()
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_TEXT_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


# The columns that hold a DTN-style numeric id. Excel treats a plain number
# typed into a cell as a real number, not text — so openpyxl hands it back
# as a Python int/float, and str()'ing a whole float tacks on a trailing
# ".0" ("20190708085612.0"). Clean that here so it never reaches the
# database (or the 14-digit check below) with the artifact still attached.
DTN_FIELD_KEYS = {"letter_dtn", "registration_dtn", "donation_reg_no"}


def _clean_dtn_cell(value):
    """Collapse a DTN-style cell back to a clean digit string — whether
    the ".0" artifact shows up as a native float/int from Excel, or as
    text that already picked it up from an earlier bad export/re-import."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


@router.get("/download-template")
async def download_donation_template():
    """Download a formatted Donation Database Excel upload template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Donation Template"

    default_fill = PatternFill("solid", start_color="FFFFFF")
    header_font = Font(name="Arial Narrow", bold=True, size=12, color="000000")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, (label, width, fill_color, _key) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = PatternFill("solid", start_color=fill_color) if fill_color else default_fill
        cell.alignment = header_align
        cell.border = thin
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 75
    ws.freeze_panes = "B2"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=donation_database_template.xlsx"},
    )


@router.get("/", response_model=List[DonationResponse])
def list_donations(
    skip: int = Query(0, ge=0),
    limit: Optional[int] = Query(None, ge=1, le=5000),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """All donation records by default (this table is currently small).
    Pass `skip`/`limit` to page through it once it grows — omitting them
    keeps today's "fetch everything, filter client-side" behavior."""
    return crud.get_all_donations(db, skip=skip, limit=limit)


@router.get("/deleted", response_model=List[DonationResponse])
def list_deleted_donations(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """Soft-deleted records — the "trash" view Delete/Restore work against.
    Must stay registered before /{donation_id} or FastAPI would try to
    parse "deleted" as an id."""
    return crud.get_deleted_donations(db)


@router.get("/{donation_id}", response_model=DonationResponse)
def get_donation(
    donation_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    donation = crud.get_donation_by_id(db, donation_id)
    if not donation:
        raise HTTPException(status_code=404, detail="Donation record not found")
    return donation


@router.post("/", response_model=DonationResponse, status_code=201)
def create_donation(
    payload: DonationCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    try:
        return crud.create_donation(db, payload, current_user.username)
    except crud.DuplicateLetterDtnError as exc:
        raise HTTPException(
            status_code=409,
            detail=f"Letter DTN \"{exc.letter_dtn}\" already exists (record #{exc.existing_id}).",
        ) from exc


@router.put("/{donation_id}", response_model=DonationResponse)
def update_donation(
    donation_id: int,
    payload: DonationUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    try:
        donation = crud.update_donation(
            db, donation_id, payload, current_user.username,
            expected_version=payload.version,
        )
    except crud.StaleVersionError as exc:
        raise HTTPException(
            status_code=409,
            detail=(
                "This record was changed by someone else since you opened it. "
                "Refresh and try again."
            ),
        ) from exc
    if not donation:
        raise HTTPException(status_code=404, detail="Donation record not found")
    return donation


@router.delete("/{donation_id}", status_code=204)
def delete_donation(
    donation_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """Soft delete — the row stays in the database (and keeps its change
    log) but disappears from the list/get/update endpoints. Undo with
    POST /{donation_id}/restore."""
    donation = crud.soft_delete_donation(db, donation_id, current_user.username)
    if not donation:
        raise HTTPException(status_code=404, detail="Donation record not found")


@router.post("/{donation_id}/restore", response_model=DonationResponse)
def restore_donation(
    donation_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    donation = crud.restore_donation(db, donation_id, current_user.username)
    if not donation:
        raise HTTPException(status_code=404, detail="Deleted donation record not found")
    return donation


@router.get("/{donation_id}/change-log", response_model=List[DonationChangeLogResponse])
def get_donation_change_log(
    donation_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    # Uses the "any" lookup (not filtered by is_deleted) so a deleted
    # record's history is still reachable — the audit trail shouldn't
    # disappear just because the record itself is hidden.
    if not crud.get_donation_by_id_any(db, donation_id):
        raise HTTPException(status_code=404, detail="Donation record not found")
    return crud.get_change_log(db, donation_id)


EXPECTED_HEADER_HINTS = ("Letter DTN", "Donor", "Product Name")


def _pick_import_sheet(wb):
    """Prefer the sheet named like our own template; otherwise use the
    first sheet by position — never `wb.active`, which reflects whatever
    tab was open when the file was last saved, not which sheet actually
    holds the data (a real risk if someone re-uploads a multi-sheet
    workbook instead of a copy of the downloaded template)."""
    for name in wb.sheetnames:
        if "donation template" in name.lower():
            return wb[name]
    return wb.worksheets[0]


@router.post("/upload-excel", response_model=DonationUploadResult)
async def upload_donation_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """Import donations from a filled-in copy of the download-template
    workbook. Columns are read strictly by position (matches
    TEMPLATE_COLUMNS order), not by header text — so the header row is
    only used here to sanity-check the shape, not to map columns."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Invalid file type. Must be .xlsx or .xls")

    contents = await file.read()
    try:
        wb = load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {exc}")

    ws = _pick_import_sheet(wb)
    field_keys = [key for (_label, _w, _fill, key) in TEMPLATE_COLUMNS]

    header_cells = next(ws.iter_rows(min_row=1, max_row=1, max_col=len(field_keys), values_only=True), ())
    header_text = " ".join(str(c) for c in header_cells if c)
    if not all(hint in header_text for hint in EXPECTED_HEADER_HINTS):
        raise HTTPException(
            status_code=400,
            detail=(
                "This file's columns don't look like the Donation Database "
                "template. Please use a filled-in copy of Download Template."
            ),
        )

    rows = []
    for row in ws.iter_rows(min_row=2, max_col=len(field_keys), values_only=True):
        row_dict = {}
        for key, value in zip(field_keys, row):
            if value is None:
                row_dict[key] = None
            elif key in DATE_FIELD_KEYS:
                row_dict[key] = _parse_import_date(value)
            elif key in DTN_FIELD_KEYS:
                row_dict[key] = _clean_dtn_cell(value)
            else:
                row_dict[key] = str(value).strip()
        rows.append(row_dict)

    created, skipped_duplicates, skipped_no_dtn, skipped_invalid_dtn, errors = crud.bulk_create_donations(
        db, rows, current_user.username
    )
    return DonationUploadResult(
        created=created,
        skipped_duplicates=skipped_duplicates,
        skipped_no_dtn=skipped_no_dtn,
        skipped_invalid_dtn=skipped_invalid_dtn,
        failed=len(errors),
        errors=errors,
    )


@router.post("/upload-preview")
async def preview_donation_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    """Read-only dry run: parses the file and reports what /upload-excel
    would do (which rows would be added, which would be skipped and why —
    missing DTN, duplicate DTN) without writing anything. Powers the
    upload modal's "Check File" confirmation step, so the user sees the
    outcome before committing to the real import. The dedup/required-DTN
    rules here mirror /upload-excel exactly so the preview never promises
    something the real upload won't do."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Invalid file type. Must be .xlsx or .xls")

    contents = await file.read()
    try:
        wb = load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {exc}")

    ws = _pick_import_sheet(wb)
    field_keys = [key for (_label, _w, _fill, key) in TEMPLATE_COLUMNS]

    header_cells = next(ws.iter_rows(min_row=1, max_row=1, max_col=len(field_keys), values_only=True), ())
    header_text = " ".join(str(c) for c in header_cells if c)
    if not all(hint in header_text for hint in EXPECTED_HEADER_HINTS):
        raise HTTPException(
            status_code=400,
            detail=(
                "This file's columns don't look like the Donation Database "
                "template. Please use a filled-in copy of Download Template."
            ),
        )

    existing_dtns = {
        dtn
        for (dtn,) in db.query(crud.Donation.letter_dtn).filter(crud.Donation.letter_dtn.isnot(None)).all()
    }
    seen_in_file = set()
    will_insert, will_skip = [], []

    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=len(field_keys), values_only=True), start=2):
        row_dict = {}
        for key, value in zip(field_keys, row):
            if value is None:
                row_dict[key] = None
            elif key in DATE_FIELD_KEYS:
                row_dict[key] = _parse_import_date(value)
            elif key in DTN_FIELD_KEYS:
                row_dict[key] = _clean_dtn_cell(value)
            else:
                row_dict[key] = str(value).strip()
        if not any((v or "").strip() for v in row_dict.values() if v is not None):
            continue  # truly empty row — nothing to report

        dtn = (row_dict.get("letter_dtn") or "").strip()
        preview_row = {
            "row_number": idx,
            "dtn": dtn or "-",
            "donor": row_dict.get("donor") or "-",
            "product_name": row_dict.get("product_name") or "-",
            "date_received": row_dict.get("date_received") or "-",
        }

        if not dtn:
            will_skip.append({**preview_row, "reason": "No Letter DTN in this row — a row needs a Letter DTN to be imported."})
            continue
        if not crud.LETTER_DTN_RE.match(dtn):
            will_skip.append({**preview_row, "reason": f"Letter DTN \"{dtn}\" isn't a 14-digit number — won't import."})
            continue
        if dtn in seen_in_file:
            will_skip.append({**preview_row, "reason": "Duplicate Letter DTN — appears more than once in this file."})
            continue
        if dtn in existing_dtns:
            will_skip.append({**preview_row, "reason": "Letter DTN already exists in the system — will be skipped, not overwritten."})
            continue

        seen_in_file.add(dtn)
        will_insert.append(preview_row)

    return {
        "total_rows": len(will_insert) + len(will_skip),
        "insert_count": len(will_insert),
        "skip_count": len(will_skip),
        "will_insert": will_insert,
        "will_skip": will_skip,
    }
