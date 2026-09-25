# FILE: app/services/clinical_trial_excel.py
import io
import re
from datetime import date, datetime, timedelta
from typing import List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.clinical_trial_columns import (
    CLINICAL_TRIAL_COLUMNS,
    CLINICAL_TRIAL_DRUG_COLUMNS,
    ALL_EXCEL_COLUMNS,
    VALID_PHASES,
)
from app.models.clinical_trial import ClinicalTrial
from app.schemas.clinical_trial import ClinicalTrialCreate, ClinicalTrialDrugCreate

SHEET_NAME = "Clinical Trials"
HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

_DATE_FORMATS = [
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%b %d, %Y",
    "%B %d, %Y",
    "%d-%b-%Y",
    "%d %B %Y",
    "%Y/%m/%d",
    "%m-%d-%Y",
]

_ARABIC_TO_ROMAN_PHASE = {"1": "I", "2": "II", "3": "III", "4": "IV"}
_PHASE_MAX_LEN = 20

_EXCEL_EPOCH = datetime(1899, 12, 30)


def _write_header(ws):
    for col_idx, column in enumerate(ALL_EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column["label"])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            18, len(column["label"]) + 2
        )


def build_template_workbook() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _write_header(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _fill_merged_cells(ws) -> None:
    merged_ranges = list(ws.merged_cells.ranges)
    for merged_range in merged_ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = ws.cell(row=min_row, column=min_col).value
        ws.unmerge_cells(str(merged_range))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col, value=top_left_value)


def _parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return (_EXCEL_EPOCH + timedelta(days=value)).date()
        except (OverflowError, ValueError):
            raise ValueError(f"Unrecognized date value: '{value}'")
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(
        f"Unrecognized date format: '{text}'. Expected YYYY-MM-DD (e.g. 2026-06-01)."
    )


def _parse_int(value) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, str):
        value = value.replace(",", "").strip()
        if not value:
            return 0
    return int(float(value))


def _clean_text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float):
        text = str(int(value)) if value.is_integer() else str(value)
    elif isinstance(value, int):
        text = str(value)
    else:
        text = str(value)
    text = text.strip()
    return text or None


def _normalize_phase(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"(?i)^phase\s*", "", text).strip()
    if not text:
        return None
    if "/" in text:
        parts = [_normalize_phase(part) for part in text.split("/")]
        parts = [part for part in parts if part]
        result = "/".join(parts) if parts else None
        return result[:_PHASE_MAX_LEN] if result else None
    match = re.match(r"^(\d+)\s*([a-zA-Z]?)$", text)
    if match:
        digit, suffix = match.groups()
        roman = _ARABIC_TO_ROMAN_PHASE.get(digit)
        if roman:
            result = f"{roman}{suffix.upper()}" if suffix else roman
            return result[:_PHASE_MAX_LEN]
    text_upper = text.upper()
    if text_upper in VALID_PHASES:
        return text_upper
    match = re.match(r"^(I{1,3}|IV)([A-Z]?)$", text_upper)
    if match:
        roman, suffix = match.groups()
        result = f"{roman}{suffix}" if suffix else roman
        return result[:_PHASE_MAX_LEN]
    return None


def _parse_row(raw_row) -> dict:
    """Parses one Excel row into {field: value} for ALL_EXCEL_COLUMNS."""
    row_data = {}
    for col_idx, column in enumerate(ALL_EXCEL_COLUMNS):
        value = raw_row[col_idx] if col_idx < len(raw_row) else None
        if column["field"] == "phase":
            value = _normalize_phase(value)
        elif column["field"] == "il_approval_date":
            value = _parse_date(value)
        elif column["field"] == "total_qty_approve":
            value = _parse_int(value)
        else:
            value = _clean_text(value)
        row_data[column["field"]] = value
    return row_data


def _trial_key(row_data: dict):
    """
    Groups consecutive rows into the same trial.

    Real trackers commonly repeat the trial identifier (Protocol Number)
    on every drug row, but only fill in the other trial-level fields
    (Study Title, Phase, Sponsor, etc.) on ONE of those rows and leave
    the rest blank. Grouping on the full set of trial fields would treat
    a blank-vs-filled Study Title as a different trial and incorrectly
    split what is really one trial into several - which is what was
    happening before this fix.

    So grouping now keys on identity fields only: Protocol Number first
    (the field trackers repeat most reliably), falling back to CT
    Reference Number when Protocol Number is blank. If a row has neither,
    it can't be grouped with anything and becomes its own trial.
    """
    protocol_no = row_data.get("protocol_no")
    ct_ref_no = row_data.get("ct_ref_no")
    if protocol_no:
        return ("protocol_no", protocol_no)
    if ct_ref_no:
        return ("ct_ref_no", ct_ref_no)
    return ("row", id(row_data))


def parse_upload_workbook(
    file_bytes: bytes,
) -> Tuple[List[Tuple[int, ClinicalTrialCreate]], List[str]]:
    """
    Returns (rows, errors). `rows` is a list of (first_excel_row_number,
    ClinicalTrialCreate) — consecutive Excel rows sharing the same
    Protocol Number (or CT Reference Number when Protocol Number is
    blank) are grouped into ONE trial with multiple `drugs` entries. See
    `_trial_key` for the exact grouping rule and why it only keys on
    those two identity fields.
    """
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    _fill_merged_cells(ws)

    header_row = [cell.value for cell in ws[1]]
    expected_labels = [c["label"] for c in ALL_EXCEL_COLUMNS]
    if header_row[: len(expected_labels)] != expected_labels:
        return [], [
            "Uploaded file does not match the official template headers. Please use Download Template."
        ]

    errors: List[str] = []
    groups: List[dict] = []  # each: {"row_idx", "trial_fields", "drugs": [dict,...]}
    current = None

    for row_idx, raw_row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if raw_row is None or all(v is None or str(v).strip() == "" for v in raw_row):
            continue

        try:
            row_data = _parse_row(raw_row)
        except Exception as exc:
            errors.append(f"Row {row_idx}: {exc}")
            continue

        trial_fields = {
            c["field"]: row_data[c["field"]] for c in CLINICAL_TRIAL_COLUMNS
        }
        drug_fields = {
            c["field"]: row_data[c["field"]] for c in CLINICAL_TRIAL_DRUG_COLUMNS
        }
        key = _trial_key(row_data)

        has_any_drug_data = any(v not in (None, "", 0) for v in drug_fields.values())

        if current is not None and current["key"] == key:
            # Same trial as the previous row - fill in any trial-level
            # field that's still blank on the group so far. Whichever row
            # in the group actually carries the Study Title/Phase/Sponsor/
            # etc. wins; a blank cell on one row no longer blanks out a
            # value already found on another row of the same trial.
            for field, value in trial_fields.items():
                if current["trial_fields"].get(field) in (None, "") and value not in (
                    None,
                    "",
                ):
                    current["trial_fields"][field] = value
            if has_any_drug_data:
                current["drugs"].append(drug_fields)
        else:
            current = {
                "row_idx": row_idx,
                "key": key,
                "trial_fields": trial_fields,
                "drugs": [drug_fields] if has_any_drug_data else [],
            }
            groups.append(current)

    rows: List[Tuple[int, ClinicalTrialCreate]] = []
    for group in groups:
        try:
            drugs = [ClinicalTrialDrugCreate(**d) for d in group["drugs"]]
            trial = ClinicalTrialCreate(**group["trial_fields"], drugs=drugs)
            rows.append((group["row_idx"], trial))
        except Exception as exc:
            errors.append(f"Row {group['row_idx']}: {exc}")

    return rows, errors


def build_export_workbook(trials: List[ClinicalTrial]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    _write_header(ws)

    row_idx = 2
    for trial in trials:
        trial_values = {
            c["field"]: getattr(trial, c["field"]) for c in CLINICAL_TRIAL_COLUMNS
        }
        drug_rows = trial.drugs or [None]
        for drug in drug_rows:
            for col_idx, column in enumerate(ALL_EXCEL_COLUMNS, start=1):
                if column["field"] in trial_values:
                    value = trial_values[column["field"]]
                else:
                    value = getattr(drug, column["field"]) if drug else None
                if isinstance(value, date):
                    value = value.isoformat()
                ws.cell(row=row_idx, column=col_idx, value=value)
            row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
